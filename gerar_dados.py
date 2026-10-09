#!/usr/bin/env python3
"""
Gera os arquivos de dados do mapa a partir da planilha base.

Uso:
    python scripts/gerar_dados.py --planilha "CAMINHO/BASE_OTIMIZACAO_ROTAS_URNAS_2026.xlsx"

Opções:
    --sem-rede      não consulta o serviço de rotas; usa só o cache e, sem cache, linha reta
    --cenario       nome do cenário a gerar (padrão: atual). Gera data/cenario-<nome>.geojson

Abas lidas da planilha:
    PONTOS_SAIDA   ID_SAIDA, MUNICIPIO, NOME_PONTO_SAIDA, ENDERECO, BAIRRO, LATITUDE, LONGITUDE, UEOP, CIA
    LOCAIS         ID_LOCAL, MUNICIPIO, ID_SAIDA, LOCAL_VOTACAO, ENDERECO, BAIRRO, LATITUDE, LONGITUDE, URNAS
    ROTAS_2026     ID_ROTA, MUNICIPIO, ROTA, ID_SAIDA, UEOP, CIA, MODALIDADE
    PARADAS_2026   ID_ROTA, ORDEM, ID_LOCAL

Saídas:
    data/cenario-<nome>.geojson   origens, destinos e rotas
    data/unidades.json            unidades (3ª RPM, BPM e Cias) e seus escudos
    scripts/cache_rotas.json      traçados já calculados (evita recalcular)

Requer: Python 3.9+ e openpyxl (pip install openpyxl).
"""
import argparse, datetime as dt, json, math, os, sys, time, unicodedata, urllib.request

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CACHE = os.path.join(RAIZ, "scripts", "cache_rotas.json")
OSRM = "https://router.project-osrm.org/route/v1/driving/{}?overview=full&geometries=polyline&steps=false"

# Campos que nunca vão para os arquivos publicados, mesmo que existam na planilha.
CAMPOS_SENSIVEIS = {"EQUIPE", "VIATURA", "PLACA", "HORARIO", "HORARIOS", "COMANDANTE", "MILITARES", "TELEFONE", "NR_PM"}

# Unidades: a Cia de cada rota vem da planilha; aqui ficam nome, escudo e hierarquia.
UNIDADES = {
    "regiao": {"id": "3rpm", "nome": "3ª RPM", "escudo": "assets/escudos/3rpm.webp"},
    "unidades": [
        {"id": "52bpm", "nome": "52º BPM", "tipo": "BPM", "pai": "3rpm", "escudo": "assets/escudos/52bpm.webp"},
        {"id": "248cia", "nome": "248ª Cia PM", "tipo": "CIA", "pai": "52bpm", "escudo": "assets/escudos/52bpm.webp"},
        {"id": "239cia", "nome": "239ª Cia PM", "tipo": "CIA", "pai": "52bpm", "escudo": "assets/escudos/52bpm.webp"},
        {"id": "250cia", "nome": "250ª Cia PM", "tipo": "CIA", "pai": "52bpm", "escudo": "assets/escudos/52bpm.webp"},
    ],
}
MODALIDADES = {"VIATURA PM": "VIATURA_PM", "EMBARCADO TRE": "EMBARCADO_TRE"}


def sem_acento(s):
    return unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().upper().strip()


def id_unidade(nome):
    n = sem_acento(nome).replace(" ", "")
    for u in UNIDADES["unidades"]:
        if sem_acento(u["nome"]).replace(" ", "") == n:
            return u["id"]
    return None


def ler_aba(wb, nome):
    if nome not in wb.sheetnames:
        sys.exit(f"Aba {nome} não encontrada na planilha.")
    ws = wb[nome]
    linhas = list(ws.iter_rows(values_only=True))
    cab = [sem_acento(c) for c in linhas[0]]
    out = []
    for r in linhas[1:]:
        if r and r[0] not in (None, ""):
            out.append({k: v for k, v in zip(cab, r) if k and k not in CAMPOS_SENSIVEIS})
    return out


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return None


def decodificar_polyline(s, precisao=5):
    coords, i, lat, lon, f = [], 0, 0, 0, 10 ** precisao
    while i < len(s):
        for eixo in (0, 1):
            res = shift = 0
            while True:
                b = ord(s[i]) - 63; i += 1
                res |= (b & 0x1F) << shift; shift += 5
                if b < 0x20:
                    break
            d = ~(res >> 1) if res & 1 else res >> 1
            if eixo == 0:
                lat += d
            else:
                lon += d
        coords.append([round(lon / f, 5), round(lat / f, 5)])
    return coords


def simplificar(coords, tol=0.00003):
    """Douglas-Peucker: remove pontos que desviam menos de ~3 m da linha (sem perda visível no mapa)."""
    if len(coords) < 3:
        return coords
    manter = [False] * len(coords)
    manter[0] = manter[-1] = True
    pilha = [(0, len(coords) - 1)]
    while pilha:
        i, j = pilha.pop()
        (x1, y1), (x2, y2) = coords[i], coords[j]
        dx, dy = x2 - x1, y2 - y1
        n = math.hypot(dx, dy) or 1e-12
        dmax, kmax = 0, None
        for k in range(i + 1, j):
            x0, y0 = coords[k]
            dd = abs(dy * x0 - dx * y0 + x2 * y1 - y2 * x1) / n
            if dd > dmax:
                dmax, kmax = dd, k
        if kmax is not None and dmax > tol:
            manter[kmax] = True
            pilha += [(i, kmax), (kmax, j)]
    return [c for c, m in zip(coords, manter) if m]


def haversine(a, b):
    r = 6371000
    p1, p2 = math.radians(a[1]), math.radians(b[1])
    dp, dl = p2 - p1, math.radians(b[0] - a[0])
    return 2 * r * math.asin(math.sqrt(math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2))


def calcular(pontos, cache, sem_rede):
    chave = ";".join(f"{lon:.6f},{lat:.6f}" for lon, lat in pontos)
    if chave in cache:
        return cache[chave], True
    if not sem_rede:
        try:
            with urllib.request.urlopen(OSRM.format(chave), timeout=30) as resp:
                j = json.loads(resp.read().decode())
            time.sleep(1.1)  # limite de uso do servidor público
            if j.get("code") == "Ok":
                ro = j["routes"][0]
                cache[chave] = {"d": ro["distance"], "t": ro["duration"], "legs": [[l["distance"], l["duration"]] for l in ro["legs"]],
                                "g": ro["geometry"], "snap": [w["distance"] for w in j["waypoints"]]}
                return cache[chave], True
        except Exception as e:  # sem rede ou serviço fora do ar
            print(f"  aviso: rota não calculada ({e})")
    return None, False


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--planilha", required=True)
    ap.add_argument("--sem-rede", action="store_true")
    ap.add_argument("--cenario", default="atual")
    a = ap.parse_args()
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("Instale o openpyxl: pip install openpyxl")
    wb = load_workbook(a.planilha, data_only=True)
    saidas = {s["ID_SAIDA"]: s for s in ler_aba(wb, "PONTOS_SAIDA")}
    locais = {l["ID_LOCAL"]: l for l in ler_aba(wb, "LOCAIS")}
    rotas = ler_aba(wb, "ROTAS_2026")
    paradas = ler_aba(wb, "PARADAS_2026")
    cache = json.load(open(CACHE, encoding="utf-8"))["rotas"] if os.path.exists(CACHE) else {}

    feats, pend = [], []
    cia_saida = {}
    for sid, s in saidas.items():
        lon, lat = num(s.get("LONGITUDE")), num(s.get("LATITUDE"))
        cia = s.get("CIA")
        cia_saida[sid] = cia
        feats.append({"type": "Feature", "geometry": {"type": "Point", "coordinates": [lon, lat]},
                      "properties": {"tipo": "origem", "id": sid, "municipio": s.get("MUNICIPIO"), "nome": s.get("NOME_PONTO_SAIDA"),
                                     "endereco": s.get("ENDERECO"), "bairro": s.get("BAIRRO"), "ueop": s.get("UEOP"),
                                     "cia": cia, "unidade": id_unidade(cia), "bpm": id_unidade(s.get("UEOP"))}})

    por_rota = {}
    for p in paradas:
        por_rota.setdefault(p["ID_ROTA"], []).append(p)
    rotas_do_local = {}
    totais = {}
    for r in rotas:
        rid = r["ID_ROTA"]
        ps = sorted(por_rota.get(rid, []), key=lambda p: num(p["ORDEM"]))
        if not ps:
            pend.append(f"{rid}: rota sem paradas"); continue
        s = saidas[r["ID_SAIDA"]]
        pts = [(num(s["LONGITUDE"]), num(s["LATITUDE"]))]
        faltou = False
        for p in ps:
            l = locais.get(p["ID_LOCAL"])
            if not l or num(l.get("LATITUDE")) is None:
                pend.append(f"{rid}: local {p['ID_LOCAL']} sem coordenada"); faltou = True; continue
            pts.append((num(l["LONGITUDE"]), num(l["LATITUDE"])))
            rotas_do_local.setdefault(p["ID_LOCAL"], []).append(rid)
        if faltou or len(pts) < 2:
            continue
        res, ok = calcular(pts, cache, a.sem_rede)
        if ok:
            geom = simplificar(decodificar_polyline(res["g"]))
            dist_km, tempo_min = round(res["d"] / 1000, 2), round(res["t"] / 60, 1)
            trechos = [{"km": round(l[0] / 1000, 2), "min": round(l[1] / 60, 1)} for l in res["legs"]]
            snap = round(max(res["snap"]))
        else:
            geom = [[round(x, 5), round(y, 5)] for x, y in pts]
            dist_km = round(sum(haversine(pts[i], pts[i + 1]) for i in range(len(pts) - 1)) / 1000, 2)
            tempo_min, trechos, snap = None, [], None
            pend.append(f"{rid}: rota não calculada (linha reta)")
        mod = MODALIDADES.get(sem_acento(r.get("MODALIDADE")), "NAO_INFORMADA")
        nomes = [locais[p["ID_LOCAL"]]["LOCAL_VOTACAO"] for p in ps if p["ID_LOCAL"] in locais]
        # urnas da parada (aba PARADAS_2026); sem essa coluna, usa as urnas do local
        urnas = sum(int(num(p.get("URNAS") if "URNAS" in p else locais.get(p["ID_LOCAL"], {}).get("URNAS")) or 0) for p in ps)
        cia = r.get("CIA") or cia_saida.get(r["ID_SAIDA"])
        feats.append({"type": "Feature", "geometry": {"type": "LineString", "coordinates": geom},
                      "properties": {"tipo": "rota", "id": rid, "rota": r.get("ROTA"), "municipio": r.get("MUNICIPIO"),
                                     "id_saida": r["ID_SAIDA"], "origem": s.get("NOME_PONTO_SAIDA"), "ueop": r.get("UEOP"),
                                     "cia": cia, "unidade": id_unidade(cia), "bpm": id_unidade(r.get("UEOP")),
                                     "modalidade": mod, "paradas": [p["ID_LOCAL"] for p in ps], "destinos": nomes,
                                     "qtd_paradas": len(ps), "urnas": urnas, "distancia_km": dist_km, "tempo_min": tempo_min,
                                     "trechos": trechos, "calculada": ok, "distancia_via_m": snap}})
        t = totais.setdefault(r.get("MUNICIPIO"), {"rotas": 0, "destinos": 0})
        t["rotas"] += 1

    for lid, l in locais.items():
        lon, lat = num(l.get("LONGITUDE")), num(l.get("LATITUDE"))
        if lon is None or lat is None:
            pend.append(f"{lid}: local sem coordenada (não plotado)"); continue
        cia = cia_saida.get(l.get("ID_SAIDA"))
        feats.append({"type": "Feature", "geometry": {"type": "Point", "coordinates": [round(lon, 7), round(lat, 7)]},
                      "properties": {"tipo": "destino", "id": lid, "municipio": l.get("MUNICIPIO"), "id_saida": l.get("ID_SAIDA"),
                                     "nome": l.get("LOCAL_VOTACAO"), "endereco": l.get("ENDERECO"), "bairro": l.get("BAIRRO"),
                                     "urnas": int(num(l.get("URNAS")) or 0), "cia": cia, "unidade": id_unidade(cia),
                                     "bpm": "52bpm", "rotas": rotas_do_local.get(lid, []), "sem_rota": lid not in rotas_do_local}})
        totais.setdefault(l.get("MUNICIPIO"), {"rotas": 0, "destinos": 0})["destinos"] += 1

    obs = [f"{f['properties']['id']}: {f['properties']['nome']} sem rota de escolta em 2026"
           for f in feats if f["properties"]["tipo"] == "destino" and f["properties"]["sem_rota"]]
    obs += [f"{f['properties']['id']}: destino a {f['properties']['distancia_via_m']} m da via mais próxima; o último trecho não aparece no traçado"
            for f in feats if f["properties"]["tipo"] == "rota" and (f["properties"]["distancia_via_m"] or 0) > 300]
    gj = {"type": "FeatureCollection",
          "metadata": {"cenario": a.cenario, "gerado_em": dt.datetime.now().strftime("%d/%m/%Y %H:%M"),
                       "fonte": "BASE_OTIMIZACAO_ROTAS_URNAS_2026.xlsx", "roteamento": "OSRM (perfil carro)",
                       "totais": totais, "pendencias": pend, "observacoes": obs},
          "features": feats}
    os.makedirs(os.path.join(RAIZ, "data"), exist_ok=True)
    with open(os.path.join(RAIZ, "data", f"cenario-{a.cenario}.geojson"), "w", encoding="utf-8") as f:
        json.dump(gj, f, ensure_ascii=False, separators=(",", ":"))
    with open(os.path.join(RAIZ, "data", "unidades.json"), "w", encoding="utf-8") as f:
        json.dump(UNIDADES, f, ensure_ascii=False, indent=1)
    with open(CACHE, "w", encoding="utf-8") as f:
        json.dump({"fonte": "router.project-osrm.org (perfil carro)", "rotas": cache}, f, ensure_ascii=False, separators=(",", ":"))

    print(f"Cenário '{a.cenario}' gerado.")
    for m, t in totais.items():
        print(f"  {m}: {t['rotas']} rotas, {t['destinos']} destinos")
    for titulo, lista in (("Pendências", pend), ("Observações", obs)):
        if lista:
            print(f"{titulo}:")
            for p in lista:
                print("  -", p)


if __name__ == "__main__":
    main()
